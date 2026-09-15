#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
import unicodedata
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

NEEDS_REVIEW = "NEEDS_REVIEW"

TARGET_FOLDERS = {
    "cerere-finantare": "cerere_finantare",
    "manual-beneficiar": "manual_beneficiar",
    "descriere-activitati": "descriere_activitati",
    "fise-post": "fisa_post",
    "raportari-aprobate-oir": "raportare_aprobata_oir",
    "livrabile-istorice": "livrabil_istoric",
}

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
UNSUPPORTED_REFERENCE_EXTENSIONS = {
    ".csv",
    ".doc",
    ".ics",
    ".jpeg",
    ".jpg",
    ".msg",
    ".odt",
    ".png",
    ".ppt",
    ".pptx",
    ".rtf",
    ".xls",
    ".xlsx",
}
DOCUMENT_EXTENSIONS = SUPPORTED_EXTENSIONS | UNSUPPORTED_REFERENCE_EXTENSIONS
EXCLUDED_DIRS = {".git", ".next", ".codex", ".amplify", "node_modules", "out", "build"}

MONTHS = {
    "IANUARIE": 1,
    "FEBRUARIE": 2,
    "MARTIE": 3,
    "APRILIE": 4,
    "MAI": 5,
    "IUNIE": 6,
    "IULIE": 7,
    "AUGUST": 8,
    "SEPT": 9,
    "SEPTEMBRIE": 9,
    "OCTOMBRIE": 10,
    "NOIEMBRIE": 11,
    "DECEMBRIE": 12,
}

METADATA_HEADERS = [
    "fileName",
    "sourceFileName",
    "sourcePath",
    "sourceType",
    "category",
    "approvalStatus",
    "expertId",
    "expertName",
    "expertRole",
    "positionInProject",
    "projectCode",
    "month",
    "year",
    "saCode",
    "activityName",
    "reviewNotes",
]


@dataclass
class SourceDocument:
    path: Path
    classification: str
    target_folder: str
    source_type: str


@dataclass
class ReportMetadata:
    fileName: str
    sourceFileName: str
    sourcePath: str
    sourceType: str
    category: str
    approvalStatus: str
    expertId: str
    expertName: str
    expertRole: str
    positionInProject: str
    projectCode: str
    month: str
    year: str
    saCode: str
    activityName: str
    reviewNotes: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Prepare rag-seed by converting supported source documents and generating metadata.csv.",
    )
    parser.add_argument("--repo", default=".", help="Repository root. Defaults to current directory.")
    parser.add_argument("--target", default="rag-seed", help="Target RAG seed directory.")
    parser.add_argument("--default-category", default="", help="Default PEO category written to metadata.csv when it cannot be inferred.")
    parser.add_argument("--default-expert-id", default="", help="Default expertId when local app data is the source of truth.")
    parser.add_argument("--default-expert-name", default="", help="Default expertName when local app data is the source of truth.")
    parser.add_argument("--default-expert-role", default="", help="Default expertRole when it cannot be extracted safely.")
    parser.add_argument("--default-position-in-project", default="", help="Default positionInProject when it cannot be extracted safely.")
    parser.add_argument("--default-project-code", default="", help="Default projectCode when it cannot be extracted safely.")
    parser.add_argument(
        "--source",
        action="append",
        default=[],
        help="Additional source directory or file. Can be passed multiple times.",
    )
    parser.add_argument("--dry-run", action="store_true", help="Plan only; do not write outputs.")
    parser.add_argument("--overwrite", action="store_true", help="Overwrite existing converted text files.")
    parser.add_argument("--max-files", type=int, default=0, help="Optional conversion cap for testing.")
    return parser.parse_args()


def normalized_token(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^A-Za-z0-9]+", " ", value).strip().upper()


def slugify(value: str, separator: str = "-") -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"[^A-Za-z0-9]+", separator, value).strip(separator).lower()
    return value or "document"


def title_case_name(tokens: list[str]) -> str:
    cleaned = [token.lower().capitalize() for token in tokens if token]
    return " ".join(cleaned) if cleaned else NEEDS_REVIEW


def relative_to_repo(path: Path, repo: Path) -> str:
    try:
        return str(path.resolve().relative_to(repo.resolve())).replace("\\", "/")
    except ValueError:
        return str(path.resolve())


def ensure_structure(target: Path, dry_run: bool) -> None:
    for folder in TARGET_FOLDERS:
        if not dry_run:
            (target / folder).mkdir(parents=True, exist_ok=True)


def iter_candidate_files(root: Path, target: Path) -> Iterable[Path]:
    if root.is_file():
        if root.suffix.lower() in DOCUMENT_EXTENSIONS:
            yield root
        return

    if not root.exists():
        return

    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.name.upper() in {"README.MD"}:
            continue
        parts = set(path.parts)
        if parts & EXCLUDED_DIRS:
            continue
        if path.name in {"metadata.csv", "conversion-report.json", "review-needed.csv", "unclassified-documents.csv"}:
            continue
        if path.suffix.lower() not in DOCUMENT_EXTENSIONS:
            continue
        yield path


def classify_file(path: Path, repo: Path, target: Path) -> SourceDocument | None:
    extension = path.suffix.lower()
    rel = relative_to_repo(path, repo)
    rel_norm = normalized_token(rel)
    rel_tokens = set(rel_norm.split())

    if extension not in DOCUMENT_EXTENSIONS:
        return None

    for folder, source_type in TARGET_FOLDERS.items():
        if f"RAG SEED {normalized_token(folder)}" in rel_norm:
            return SourceDocument(path, folder, folder, source_type)

    if "RAPORTARI PA" in rel_norm or parse_report_identity(path).get("hasMonthYear"):
        return SourceDocument(path, "raportare_aprobata_oir", "raportari-aprobate-oir", "raportare_aprobata_oir")
    if "LIVRABILE ISTORICE" in rel_norm or "LIVRABIL" in rel_norm:
        return SourceDocument(path, "livrabil_istoric", "livrabile-istorice", "livrabil_istoric")
    if "CERERE" in rel_tokens and "FINANTARE" in rel_tokens:
        return SourceDocument(path, "cerere_finantare", "cerere-finantare", "cerere_finantare")
    if {"MANUAL", "BENEFICIAR"} & rel_tokens or ("MANUALUL" in rel_tokens and "BENEFICIARULUI" in rel_tokens):
        return SourceDocument(path, "manual_beneficiar", "manual-beneficiar", "manual_beneficiar")
    if "DESCRIERE ACTIVITATI" in rel_norm or "DESCRIERE ACTIVITATE" in rel_norm or "SERVICII PA" in rel_norm:
        return SourceDocument(path, "descriere_activitati", "descriere-activitati", "descriere_activitati")
    if "FISE POST" in rel_norm or "FISA POST" in rel_norm:
        return SourceDocument(path, "fisa_post", "fise-post", "fisa_post")
    if "RAG SEED" in rel_norm:
        return SourceDocument(path, "unclassified", "needs-review", "other")
    return None


def canonical_report_stem(path: Path) -> str:
    stem = path.stem
    stem = re.sub(r"(?i)ocr$", "", stem)
    stem = re.sub(r"(?i)[_\-\s]*ocr[_\-\s]*", "_", stem)
    stem = stem.replace("!", "")
    stem = re.sub(r"(20[0-9]{2})[_\-\s]+[0-9]+$", r"\1", stem)
    return stem.strip("_- ")


def parse_report_identity(path: Path) -> dict[str, object]:
    stem = canonical_report_stem(path)
    tokens = [token for token in re.split(r"[_\s\-]+", normalized_token(stem)) if token]
    month_index = None
    month = None
    for index, token in enumerate(tokens):
        if token in MONTHS:
            month_index = index
            month = MONTHS[token]
            break

    year = None
    if month_index is not None:
        for token in tokens[month_index + 1 :]:
            if re.fullmatch(r"20[0-9]{2}", token):
                year = int(token)
                break

    name_start = 0
    prefix_tokens = {"ANEXA", "RA", "RAPORT", "ACTIVITATE", "DE"}
    if tokens and re.fullmatch(r"[0-9]{4,}", tokens[0]):
        name_start = 1
    for index, token in enumerate(tokens[: month_index or len(tokens)]):
        if token == "RA":
            name_start = index + 1
    name_tokens = tokens[name_start:month_index] if month_index is not None else tokens[name_start:]
    name_tokens = [token for token in name_tokens if token not in prefix_tokens and not token.isdigit()]
    expert_name = title_case_name(name_tokens)
    expert_id = slugify(expert_name) if expert_name != NEEDS_REVIEW else NEEDS_REVIEW

    return {
        "hasMonthYear": month is not None and year is not None,
        "month": month,
        "year": year,
        "expertName": expert_name,
        "expertId": expert_id,
        "projectCode": tokens[0] if tokens and re.fullmatch(r"[0-9]{4,}", tokens[0]) else NEEDS_REVIEW,
        "canonicalKey": f"{expert_id}:{year or NEEDS_REVIEW}:{month or NEEDS_REVIEW}",
    }


def choose_report_sources(sources: list[SourceDocument]) -> tuple[list[SourceDocument], list[dict[str, str]]]:
    groups: dict[str, list[SourceDocument]] = {}
    passthrough: list[SourceDocument] = []

    for source in sources:
        if source.target_folder != "raportari-aprobate-oir":
            passthrough.append(source)
            continue
        identity = parse_report_identity(source.path)
        key = str(identity["canonicalKey"])
        groups.setdefault(key, []).append(source)

    selected = passthrough[:]
    skipped: list[dict[str, str]] = []

    for key, group in groups.items():
        ranked = sorted(
            group,
            key=lambda item: (
                "OCR" not in item.path.stem.upper(),
                "!" in item.path.name,
                len(item.path.name),
                item.path.name.lower(),
            ),
        )
        winner = ranked[0]
        selected.append(winner)
        for duplicate in ranked[1:]:
            skipped.append(
                {
                    "fileName": duplicate.path.name,
                    "selectedInstead": winner.path.name,
                    "reason": f"duplicate_report_variant:{key}",
                },
            )

    return selected, skipped


def normalize_text(value: str) -> str:
    value = value.replace("\x00", " ")
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{4,}", "\n\n\n", value)
    return value.strip()


def extract_pdf_text(path: Path) -> tuple[str, str]:
    errors: list[str] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = normalize_text("\n\n".join(pages))
        if len(text) >= 100:
            return text, "pypdf"
    except Exception as exc:  # pragma: no cover - defensive utility script
        errors.append(f"pypdf:{exc.__class__.__name__}")

    try:
        import pdfplumber

        pages = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        text = normalize_text("\n\n".join(pages))
        method = "pdfplumber" if not errors else "pdfplumber_after_" + "_".join(errors)
        return text, method
    except Exception as exc:  # pragma: no cover - defensive utility script
        errors.append(f"pdfplumber:{exc.__class__.__name__}")

    return "", "failed:" + ";".join(errors)


def extract_docx_text(path: Path) -> tuple[str, str]:
    try:
        from docx import Document

        doc = Document(str(path))
        parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return normalize_text("\n".join(parts)), "python-docx"
    except Exception as exc:  # pragma: no cover - defensive utility script
        return "", f"failed:python-docx:{exc.__class__.__name__}"


def clean_metadata_value(value: str) -> str:
    value = re.sub(r"\s+", " ", value or "").strip(" :;\t\r\n")
    return value if value else NEEDS_REVIEW


def parse_month_year_value(value: str) -> tuple[int | None, int | None]:
    tokens = [token for token in re.split(r"[\s_\-]+", normalized_token(value)) if token]
    month = None
    year = None
    for token in tokens:
        if month is None and token in MONTHS:
            month = MONTHS[token]
        if year is None and re.fullmatch(r"20[0-9]{2}", token):
            year = int(token)
    return month, year


def value_after_label(line: str, required_tokens: set[str]) -> str | None:
    normalized_words = set(normalized_token(line).split())
    if not required_tokens.issubset(normalized_words):
        return None
    if ":" in line:
        return clean_metadata_value(line.split(":", 1)[1])
    if "|" in line:
        return clean_metadata_value(line.rsplit("|", 1)[1])
    return None


def clean_project_code(value: str) -> str:
    match = re.search(r"\b[0-9]{4,}\b", value)
    return match.group(0) if match else clean_metadata_value(value)


def extract_report_metadata_from_text(text: str) -> dict[str, str | int]:
    metadata: dict[str, str | int] = {}
    head = text[:5000]
    lines = head.splitlines()

    for line in lines[:30]:
        value = value_after_label(line, {"CODUL", "PROIECTULUI"})
        if value:
            metadata["projectCode"] = clean_project_code(value)
            break

    for line in lines[:30]:
        value = value_after_label(line, {"NUMELE", "EXPERTULUI"})
        if value:
            metadata["expertName"] = value
            break

    for line in lines[:30]:
        value = value_after_label(line, {"POZITIA", "CADRUL", "PROIECTULUI"})
        if value:
            metadata["expertRole"] = value
            break

    for line in lines[:20]:
        month, year = parse_month_year_value(line)
        if month and year:
            metadata.setdefault("month", month)
            metadata.setdefault("year", year)
            break

    return metadata

    patterns = {
        "projectCode": r"Codul proiectului:\s*([^\n\r]+)",
        "expertName": r"Numele(?:\s+și\s+prenumele)?\s+expertului:\s*([^\n\r]+)",
        "expertRole": r"Poziția\s+în\s+cadrul\s+proiectului:\s*([^\n\r]+)",
    }
    for key, pattern in patterns.items():
        match = re.search(pattern, head, flags=re.IGNORECASE)
        if match:
            metadata[key] = clean_metadata_value(match.group(1))

    for line in head.splitlines()[:20]:
        month, year = parse_month_year_value(line)
        if month and year:
            metadata.setdefault("month", month)
            metadata.setdefault("year", year)
            break

    return metadata


def read_text_source(path: Path) -> tuple[str, str]:
    data = path.read_bytes()
    encodings = ["utf-8-sig", "utf-16"]
    if data[1:200:2].count(0) > 20:
        encodings.extend(["utf-16-le", "utf-16-be"])
    encodings.append("cp1250")
    for encoding in encodings:
        try:
            return normalize_text(data.decode(encoding)), encoding
        except UnicodeDecodeError:
            continue
    return normalize_text(data.decode("utf-8", errors="replace")), "utf-8-replace"


def convert_to_text(path: Path) -> tuple[str, str]:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return extract_pdf_text(path)
    if extension == ".docx":
        return extract_docx_text(path)
    if extension in {".txt", ".md"}:
        return read_text_source(path)
    return "", f"unsupported:{extension}"


def output_file_for(source: SourceDocument, target: Path, text: str | None = None) -> Path:
    existing_target_folder = target / source.target_folder
    if source.path.suffix.lower() in {".txt", ".md"}:
        try:
            source.path.resolve().relative_to(existing_target_folder.resolve())
            return source.path
        except ValueError:
            pass

    if source.target_folder == "raportari-aprobate-oir":
        identity = parse_report_identity(source.path)
        content = extract_report_metadata_from_text(text or "") if text else {}
        expert_name = str(content.get("expertName") or identity["expertName"])
        expert_id = slugify(expert_name) if expert_name != NEEDS_REVIEW else str(identity["expertId"])
        year = content.get("year") or identity["year"] or NEEDS_REVIEW
        month = content.get("month") or identity["month"] or NEEDS_REVIEW
        if isinstance(month, int):
            month_part = f"{month:02d}"
        else:
            month_part = str(month)
        return target / source.target_folder / f"{expert_id}_{year}-{month_part}.txt"

    source_stem_norm = normalized_token(source.path.stem)
    if source.target_folder == "cerere-finantare" and "CERERE" in source_stem_norm and "FINANTARE" in source_stem_norm:
        return target / source.target_folder / "Text_Cerere_de_Finantare.txt"

    source_stem_tokens = set(source_stem_norm.split())
    if source.target_folder == "manual-beneficiar" and ({"MANUAL", "MANUALUL"} & source_stem_tokens) and ({"BENEFICIAR", "BENEFICIARULUI"} & source_stem_tokens):
        return target / source.target_folder / "Manualul Beneficiarului_in vigoare_V5.txt"

    if source.target_folder == "descriere-activitati" and "SERVICII" in source_stem_norm and "PA" in source_stem_norm:
        return target / source.target_folder / "Servicii_PA.txt"

    if source.target_folder == "livrabile-istorice":
        source_context = "_".join(source.path.with_suffix("").parts[-4:])
        safe_name = slugify(source_context, separator="_")[:90].rstrip("_")
        path_hash = hashlib.sha1(str(source.path.resolve()).encode("utf-8")).hexdigest()[:8]
        return target / source.target_folder / f"{safe_name}_{path_hash}.txt"

    safe_name = slugify(source.path.stem, separator="_")
    extension = ".md" if source.path.suffix.lower() == ".md" else ".txt"
    return target / source.target_folder / f"{safe_name}{extension}"


def file_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def write_text_if_needed(path: Path, text: str, overwrite: bool, dry_run: bool) -> str:
    if dry_run:
        return "planned"
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and not overwrite:
        existing = path.read_text(encoding="utf-8", errors="replace")
        if file_hash(existing) == file_hash(text):
            return "unchanged"
        return "exists_kept"
    path.write_text(text + "\n", encoding="utf-8")
    return "written"


def build_report_metadata(
    source: SourceDocument,
    output_path: Path,
    repo: Path,
    text: str,
    method: str,
    default_category: str,
    default_expert_id: str,
    default_expert_name: str,
    default_expert_role: str,
    default_position_in_project: str,
    default_project_code: str,
) -> ReportMetadata:
    identity = parse_report_identity(source.path)
    content = extract_report_metadata_from_text(text)
    month = content.get("month") or identity["month"] or NEEDS_REVIEW
    year = content.get("year") or identity["year"] or NEEDS_REVIEW
    expert_name = str(default_expert_name or content.get("expertName") or identity["expertName"])
    expert_id = default_expert_id or (slugify(expert_name) if expert_name != NEEDS_REVIEW else str(identity["expertId"]))
    project_code = str(content.get("projectCode") or default_project_code or identity["projectCode"] or NEEDS_REVIEW)
    expert_role = str(content.get("expertRole") or default_expert_role or NEEDS_REVIEW)
    position_in_project = str(content.get("expertRole") or default_position_in_project or default_expert_role or NEEDS_REVIEW)

    review_fields = ["expertRole", "positionInProject", "projectCode", "saCode", "activityName"]
    if expert_name == NEEDS_REVIEW:
        review_fields.append("expertName")
    if month == NEEDS_REVIEW:
        review_fields.append("month")
    if year == NEEDS_REVIEW:
        review_fields.append("year")
    if expert_role != NEEDS_REVIEW and "expertRole" in review_fields:
        review_fields.remove("expertRole")
    if position_in_project != NEEDS_REVIEW and "positionInProject" in review_fields:
        review_fields.remove("positionInProject")
    if project_code != NEEDS_REVIEW and "projectCode" in review_fields:
        review_fields.remove("projectCode")
    if len(text) < 100:
        review_fields.append("textExtraction")

    filename_month = identity["month"]
    filename_year = identity["year"]
    if content.get("month") and filename_month and content.get("month") != filename_month:
        review_fields.append("filenameMonthMismatch")
    if content.get("year") and filename_year and content.get("year") != filename_year:
        review_fields.append("filenameYearMismatch")

    review_notes = "NEEDS_REVIEW: " + ", ".join(dict.fromkeys(review_fields))
    if method.startswith("failed") or method.startswith("unsupported"):
        review_notes += f"; extraction={method}"

    return ReportMetadata(
        fileName=output_path.name,
        sourceFileName=source.path.name,
        sourcePath=relative_to_repo(source.path, repo),
        sourceType=source.source_type,
        category=default_category or NEEDS_REVIEW,
        approvalStatus="approved_oir",
        expertId=expert_id,
        expertName=expert_name,
        expertRole=expert_role,
        positionInProject=position_in_project,
        projectCode=project_code,
        month=str(month),
        year=str(year),
        saCode=NEEDS_REVIEW,
        activityName=NEEDS_REVIEW,
        reviewNotes=review_notes,
    )


def write_csv(path: Path, rows: list[dict[str, str]], headers: list[str], dry_run: bool) -> None:
    if dry_run:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_json(path: Path, payload: object, dry_run: bool) -> None:
    if dry_run:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def load_ignored_documents(target: Path, repo: Path) -> dict[str, set[str]]:
    ignored_path = target / "ignored-documents.csv"
    if not ignored_path.exists():
        return {"sources": set(), "targets": set(), "targetNames": set()}
    with ignored_path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = csv.DictReader(handle)
        sources = {row["sourcePath"].replace("\\", "/") for row in rows if row.get("sourcePath")}
        targets = {row["targetPath"].replace("\\", "/") for row in rows if row.get("targetPath")}
        target_names = {Path(target).name for target in targets}
        return {"sources": sources, "targets": targets, "targetNames": target_names}


def main() -> int:
    args = parse_args()
    repo = Path(args.repo).resolve()
    target = (repo / args.target).resolve() if not Path(args.target).is_absolute() else Path(args.target).resolve()

    default_sources = [repo]
    explicit_sources = [Path(value).resolve() for value in args.source]
    source_roots = explicit_sources or default_sources

    ensure_structure(target, args.dry_run)

    raw_sources: dict[Path, SourceDocument] = {}
    unclassified: list[dict[str, str]] = []
    unsupported: list[dict[str, str]] = []
    ignored_documents = load_ignored_documents(target, repo)
    ignored_sources = ignored_documents["sources"]
    ignored_targets = ignored_documents["targets"]
    ignored_target_names = ignored_documents["targetNames"]
    ignored_rows: list[dict[str, str]] = []

    for root in source_roots:
        for path in iter_candidate_files(root, target):
            resolved = path.resolve()
            if resolved in raw_sources:
                continue
            source_path = relative_to_repo(resolved, repo)
            if source_path in ignored_sources:
                ignored_rows.append({"fileName": resolved.name, "sourcePath": source_path, "targetPath": "", "reason": "ignored_documents_csv", "method": "", "chars": ""})
                continue
            if target in resolved.parents and resolved.parent.name == "raportari-aprobate-oir":
                continue
            source = classify_file(resolved, repo, target)
            if source is None:
                if "rag-seed" in [part.lower() for part in resolved.parts]:
                    unclassified.append({"fileName": resolved.name, "sourcePath": relative_to_repo(resolved, repo), "reason": "no_known_pa_seed_classification"})
                continue
            planned_output = relative_to_repo(output_file_for(source, target), repo)
            if planned_output in ignored_targets or Path(planned_output).name in ignored_target_names:
                ignored_rows.append({"fileName": resolved.name, "sourcePath": source_path, "targetPath": planned_output, "reason": "ignored_documents_csv_target", "method": "", "chars": ""})
                continue
            if resolved.suffix.lower() not in SUPPORTED_EXTENSIONS:
                unsupported.append(
                    {
                        "fileName": resolved.name,
                        "sourcePath": relative_to_repo(resolved, repo),
                        "classification": source.classification,
                        "reason": f"unsupported_extension:{resolved.suffix.lower()}",
                    },
                )
                continue
            raw_sources[resolved] = source

    selected_sources, duplicate_skips = choose_report_sources(list(raw_sources.values()))
    if args.max_files > 0:
        selected_sources = selected_sources[: args.max_files]

    conversions: list[dict[str, object]] = []
    report_rows: list[ReportMetadata] = []

    for index, source in enumerate(sorted(selected_sources, key=lambda item: relative_to_repo(item.path, repo).lower()), start=1):
        if args.dry_run:
            text, method = "", "dry-run"
        else:
            text, method = convert_to_text(source.path)
        output_path = output_file_for(source, target, text)
        if not args.dry_run and source.target_folder == "livrabile-istorice" and len(normalize_text(text)) < 100:
            if output_path.exists():
                output_path.unlink()
            ignored_rows.append(
                {
                    "fileName": source.path.name,
                    "sourcePath": relative_to_repo(source.path, repo),
                    "targetPath": relative_to_repo(output_path, repo),
                    "reason": "text_too_short_after_extraction",
                    "method": method,
                    "chars": str(len(text)),
                },
            )
            continue
        status = write_text_if_needed(output_path, text, args.overwrite, args.dry_run)
        conversion = {
            "index": index,
            "sourcePath": relative_to_repo(source.path, repo),
            "targetPath": relative_to_repo(output_path, repo),
            "classification": source.classification,
            "method": method,
            "status": status,
            "chars": len(text),
        }
        if len(text) < 100:
            conversion["review"] = "NEEDS_REVIEW:textExtraction"
        conversions.append(conversion)

        if source.target_folder == "raportari-aprobate-oir":
            report_rows.append(build_report_metadata(
                source,
                output_path,
                repo,
                text,
                method,
                args.default_category,
                args.default_expert_id,
                args.default_expert_name,
                args.default_expert_role,
                args.default_position_in_project,
                args.default_project_code,
            ))

    metadata_path = target / "raportari-aprobate-oir" / "metadata.csv"
    if report_rows or not metadata_path.exists():
        write_csv(metadata_path, [asdict(row) for row in sorted(report_rows, key=lambda row: (row.expertName, row.year, row.month, row.fileName))], METADATA_HEADERS, args.dry_run)

    review_rows = []
    for row in report_rows:
        fields = [header for header in METADATA_HEADERS if getattr(row, header) == NEEDS_REVIEW]
        if fields:
            review_rows.append(
                {
                    "fileName": row.fileName,
                    "sourceFileName": row.sourceFileName,
                    "fields": ";".join(fields),
                    "reviewNotes": row.reviewNotes,
                },
            )
    if report_rows or not (target / "review-needed.csv").exists():
        write_csv(target / "review-needed.csv", review_rows, ["fileName", "sourceFileName", "fields", "reviewNotes"], args.dry_run)
    write_csv(target / "unclassified-documents.csv", unclassified, ["fileName", "sourcePath", "reason"], args.dry_run)
    write_csv(target / "unsupported-documents.csv", unsupported, ["fileName", "sourcePath", "classification", "reason"], args.dry_run)
    if ignored_rows and not args.dry_run:
        ignored_path = target / "ignored-documents.csv"
        existing_ignored: list[dict[str, str]] = []
        if ignored_path.exists():
            with ignored_path.open("r", encoding="utf-8-sig", newline="") as handle:
                existing_ignored = list(csv.DictReader(handle))
        merged_ignored: dict[str, dict[str, str]] = {}
        persistent_ignored_rows = [
            row
            for row in existing_ignored + ignored_rows
            if row.get("reason") == "text_too_short_after_extraction"
        ]
        for row in persistent_ignored_rows:
            key = row.get("targetPath") or row.get("sourcePath") or row.get("fileName")
            if key:
                merged_ignored[key] = {
                    "sourcePath": row.get("sourcePath", ""),
                    "targetPath": row.get("targetPath", ""),
                    "reason": row.get("reason", ""),
                    "method": row.get("method", ""),
                    "chars": row.get("chars", ""),
                }
        write_csv(
            ignored_path,
            sorted(merged_ignored.values(), key=lambda row: (row.get("sourcePath", ""), row.get("targetPath", ""))),
            ["sourcePath", "targetPath", "reason", "method", "chars"],
            args.dry_run,
        )

    readme = (
        "# RAG seed\n\n"
        "Generated by `scripts/prepare-rag-seed.py`.\n\n"
        "- Converted reports are in `raportari-aprobate-oir/`.\n"
        "- `metadata.csv` is conservative: uncertain fields are marked `NEEDS_REVIEW`.\n"
        "- `review-needed.csv` lists fields that need manual completion.\n"
        "- `unclassified-documents.csv` lists source documents that could not be classified.\n"
        "- `unsupported-documents.csv` lists document-like files that were not converted.\n"
    )
    if not args.dry_run:
        (target / "README.md").write_text(readme, encoding="utf-8")

    report = {
        "dryRun": args.dry_run,
        "target": str(target),
        "sourceRoots": [str(root) for root in source_roots],
        "convertedOrChecked": len(conversions),
        "metadataRows": len(report_rows),
        "reviewRows": len(review_rows),
        "unclassifiedRows": len(unclassified),
        "unsupportedRows": len(unsupported),
        "ignoredRows": len(ignored_rows),
        "duplicateReportVariantsSkipped": duplicate_skips,
        "conversions": conversions,
        "ignored": ignored_rows,
    }
    write_json(target / "conversion-report.json", report, args.dry_run)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
